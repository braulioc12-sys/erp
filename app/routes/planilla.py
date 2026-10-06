"""Planilla del personal (6 oct, pedido de Braulio: AFP, ONP, 5ta categoría,
CTS, gratificación, EsSalud, vacaciones según régimen, asistencia, descuentos,
préstamos, adelantos, reconocimiento de deuda...).

Esta primera parte trae lo que la planilla necesita antes de calcular:

- Parámetros de planilla (UIT, RMV, tasas AFP/ONP/EsSalud...): tabla
  `payroll_params`, corregibles desde pantalla (ver app/payroll_params.py).
- Periodos mensuales (6 oct, fase 3b): se abre el mes, se cargan comisiones,
  bonos y descuentos, se CALCULA (sueldo proporcional a los días, descuentos
  por tardanzas/faltas/subsidio, asignación familiar, gratificación de
  julio y diciembre con su bonificación, ONP/AFP, retención de 5ta
  categoría, préstamos/adelantos/deuda reconocida y aporte EsSalud del
  empleador), se imprime la boleta de cada persona y se CIERRA: al cerrar
  quedan fijos los números, se registran las cuotas descontadas y se crean
  los pagos de planilla (para el archivo de Telecrédito). Ver
  app/payroll_calc.py para las reglas y sus límites.
- CTS (mayo y noviembre): cálculo por semestre para depositar.
- Préstamos, adelantos y reconocimientos de deuda con cuotas.
- Perfil de planilla de cada persona: si entra en planilla, régimen laboral
  (general / MYPE micro / MYPE pequeña), sistema de pensiones (ONP / AFP y
  cuál), sueldo básico, asignación familiar. Cambiar el régimen ajusta los
  días de vacaciones por año de la persona (30 general, 15 MYPE).
"""
import io
import re
from datetime import date

from flask import (
    Blueprint, Response, abort, flash, g, redirect, render_template, request, send_from_directory, url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app import storage
from app.audit import log_activity
from app.auth import permission_required, validate_csrf
from app.bulk_import import XLSX_MIME
from app.db import execute, query_all, query_one
from app.helpers import now_str, parse_float, today_str
from app.payroll_calc import cts_for, compute_line, loan_installments, loan_remaining, r2
from app.payroll_params import AFP_NAMES, get_param_rows, get_params, set_param
from app.routes.legajo import current_contract
from app.routes.viajes import _save_binary_attachment

bp = Blueprint("planilla", __name__, url_prefix="/planilla")

REGIMES = [
    ("GENERAL", "Régimen general"),
    ("MYPE_MICRO", "MYPE — Microempresa"),
    ("MYPE_PEQUENA", "MYPE — Pequeña empresa"),
]
REGIME_LABELS = dict(REGIMES)
REGIME_VACATION_DAYS = {"GENERAL": 30.0, "MYPE_MICRO": 15.0, "MYPE_PEQUENA": 15.0}
PENSION_SYSTEMS = [("ONP", "ONP"), ("AFP", "AFP"), ("NINGUNO", "Sin descuento de pensión")]
PENSION_LABELS = dict(PENSION_SYSTEMS)
AFP_LABELS = dict(AFP_NAMES)
COMMISSION_TYPES = [("FLUJO", "Comisión sobre la remuneración (flujo)"), ("MIXTA", "Comisión mixta (0% sobre la remuneración)")]


def _get_staff_or_404(staff_id):
    staff = query_one("SELECT * FROM staff WHERE id = ?", (staff_id,))
    if staff is None:
        abort(404)
    return staff


def effective_salary(staff):
    """Sueldo básico de la persona: el del perfil de planilla, o el del
    contrato vigente del legajo si no se escribió uno."""
    if staff["basic_salary"] is not None:
        return float(staff["basic_salary"])
    contract = current_contract(staff["id"])
    if contract and contract["salary"]:
        return float(contract["salary"])
    return 0.0


MONTHS_ES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
_PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
LOAN_KINDS = [
    ("PRESTAMO", "Préstamo"),
    ("ADELANTO", "Adelanto de sueldo"),
    ("RECONOCIMIENTO_DEUDA", "Reconocimiento de deuda"),
]
LOAN_KIND_LABELS = dict(LOAN_KINDS)


def period_label(period):
    return f"{MONTHS_ES[int(period[5:7])].capitalize()} {period[:4]}"


def _period_or_404(period):
    if not _PERIOD_RE.match(period or ""):
        abort(404)
    row = query_one("SELECT * FROM payroll_periods WHERE period = ?", (period,))
    if row is None:
        abort(404)
    return row


@bp.route("")
@permission_required("pagos_personal", "view")
def index():
    in_payroll = query_one("SELECT COUNT(*) AS n FROM staff WHERE status = 'ACTIVO' AND in_payroll = 1")["n"]
    periods = query_all(
        """SELECT p.*, COUNT(l.id) AS people, COALESCE(SUM(l.net_pay), 0) AS net, COALESCE(SUM(l.employer_cost), 0) AS cost
           FROM payroll_periods p LEFT JOIN payroll_lines l ON l.period_id = p.id
           GROUP BY p.id ORDER BY p.period DESC"""
    )
    return render_template(
        "planilla/index.html", in_payroll=in_payroll, periods=periods, label=period_label, default_period=today_str()[:7],
    )


@bp.route("/periodos/abrir", methods=["POST"])
@permission_required("pagos_personal", "edit")
def period_open():
    if not validate_csrf():
        abort(400)
    period = (request.form.get("period") or "").strip()
    if not _PERIOD_RE.match(period):
        flash("Elige el mes de la planilla.", "error")
        return redirect(url_for("planilla.index"))
    if query_one("SELECT id FROM payroll_periods WHERE period = ?", (period,)):
        flash("Ese mes ya está abierto.", "info")
        return redirect(url_for("planilla.period_view", period=period))
    user = getattr(g, "user", None)
    pid = execute("INSERT INTO payroll_periods (period, created_by) VALUES (?, ?)", (period, user["id"] if user else None))
    log_activity("pagos_personal", "CREAR", f"Planilla {period_label(period)}: periodo abierto", entity_type="planilla", entity_id=pid,
                 entity_url=url_for("planilla.period_view", period=period))
    flash(f"Planilla de {period_label(period)} abierta. Carga las incidencias y los conceptos del mes y luego calcula.", "success")
    return redirect(url_for("planilla.period_view", period=period))


def _lines_for(period_id):
    return query_all(
        """SELECT l.*, s.name AS staff_name, s.document_number, s.position FROM payroll_lines l
           JOIN staff s ON s.id = l.staff_id WHERE l.period_id = ? ORDER BY s.name""",
        (period_id,),
    )


@bp.route("/periodo/<period>")
@permission_required("pagos_personal", "view")
def period_view(period):
    p = _period_or_404(period)
    lines = _lines_for(p["id"])
    items = query_all(
        """SELECT i.*, s.name AS staff_name FROM payroll_items i JOIN staff s ON s.id = i.staff_id
           WHERE i.period_id = ? ORDER BY s.name, i.id""",
        (p["id"],),
    )
    totals = {k: r2(sum(l[k] for l in lines)) for k in (
        "gross_total", "pension_deduction", "fifth_deduction", "loan_deduction", "advance_deduction", "debt_deduction",
        "absence_deduction", "tardiness_deduction", "subsidy_deduction", "other_deduction", "total_deductions", "net_pay",
        "essalud_employer", "employer_cost", "gratification", "gratification_bonus",
    )}
    people = query_all("SELECT id, name FROM staff WHERE in_payroll = 1 AND status = 'ACTIVO' ORDER BY name")
    return render_template(
        "planilla/periodo.html", p=p, lines=lines, items=items, totals=totals, people=people, title=period_label(period),
    )


def _calculate_period(p):
    """Recalcula todas las boletas del periodo (solo si está ABIERTO)."""
    params = get_params()
    first = f"{p['period']}-01"
    staff_rows = query_all(
        """SELECT * FROM staff WHERE in_payroll = 1 AND (status = 'ACTIVO' OR COALESCE(termination_date, '') >= ?)
           ORDER BY name""",
        (first,),
    )
    all_items = query_all("SELECT * FROM payroll_items WHERE period_id = ?", (p["id"],))
    execute("DELETE FROM payroll_lines WHERE period_id = ?", (p["id"],))
    count = 0
    for st in staff_rows:
        mine = [i for i in all_items if i["staff_id"] == st["id"]]
        line = compute_line(st, p["period"], params, mine, effective_salary, period_id=p["id"])
        if line is None:
            continue
        cols = ["period_id", "staff_id"] + list(line.keys())
        vals = [p["id"], st["id"]] + list(line.values())
        execute(
            f"INSERT INTO payroll_lines ({', '.join(cols)}) VALUES ({', '.join('?' for _ in cols)})", tuple(vals)
        )
        count += 1
    return count


@bp.route("/periodo/<period>/calcular", methods=["POST"])
@permission_required("pagos_personal", "edit")
def period_calc(period):
    if not validate_csrf():
        abort(400)
    p = _period_or_404(period)
    if p["status"] != "ABIERTO":
        flash("Esta planilla ya está cerrada. Reábrela para recalcular.", "error")
        return redirect(url_for("planilla.period_view", period=period))
    n = _calculate_period(p)
    if n == 0:
        flash("No hay personas para calcular: marca a quién entra en planilla en el Perfil de planilla.", "error")
    else:
        warns = query_one("SELECT COUNT(*) AS n FROM payroll_lines WHERE period_id = ? AND warnings IS NOT NULL", (p["id"],))["n"]
        flash(f"Planilla calculada: {n} boleta(s)." + (f" {warns} con avisos para revisar." if warns else ""), "success")
        log_activity("pagos_personal", "CALCULAR", f"Planilla {period_label(period)}: calculada ({n} boletas)",
                     entity_type="planilla", entity_id=p["id"], entity_url=url_for("planilla.period_view", period=period))
    return redirect(url_for("planilla.period_view", period=period))


@bp.route("/periodo/<period>/cerrar", methods=["POST"])
@permission_required("pagos_personal", "edit")
def period_close(period):
    if not validate_csrf():
        abort(400)
    p = _period_or_404(period)
    back = redirect(url_for("planilla.period_view", period=period))
    if p["status"] != "ABIERTO":
        flash("Esta planilla ya está cerrada.", "info")
        return back
    earlier = query_one("SELECT period FROM payroll_periods WHERE period < ? AND status = 'ABIERTO' ORDER BY period LIMIT 1", (period,))
    if earlier:
        flash(f"Cierra primero la planilla de {period_label(earlier['period'])} (se cierran en orden).", "error")
        return back
    n = _calculate_period(p)  # el cierre siempre usa los datos más recientes
    if n == 0:
        flash("No hay boletas para cerrar.", "error")
        return back
    lines = _lines_for(p["id"])
    negative = [l["staff_name"] for l in lines if l["net_pay"] < -0.005]
    if negative:
        flash("No se puede cerrar: neto negativo en " + ", ".join(negative) + ". Ajusta préstamos, adelantos o descuentos.", "error")
        return back
    user = getattr(g, "user", None)
    created = existing = 0
    for l in lines:
        for loan, amount in loan_installments(l["staff_id"], period, exclude_period_id=p["id"]):
            execute("INSERT INTO staff_loan_payments (loan_id, period_id, amount) VALUES (?, ?, ?)", (loan["id"], p["id"], amount))
        if l["net_pay"] <= 0:
            continue
        if query_one("SELECT id FROM staff_payments WHERE staff_id = ? AND payment_type = 'PLANILLA' AND period = ?", (l["staff_id"], period)):
            existing += 1
            continue
        execute(
            """INSERT INTO staff_payments (staff_id, payment_type, period, amount, concept, status)
               VALUES (?, 'PLANILLA', ?, ?, ?, 'PENDIENTE')""",
            (l["staff_id"], period, l["net_pay"], f"Planilla {period_label(period)}"),
        )
        created += 1
    execute("UPDATE payroll_periods SET status = 'CERRADO', closed_at = ?, closed_by = ? WHERE id = ?",
            (now_str(), user["id"] if user else None, p["id"]))
    log_activity("pagos_personal", "CERRAR", f"Planilla {period_label(period)}: cerrada ({n} boletas, {created} pagos generados)",
                 entity_type="planilla", entity_id=p["id"], entity_url=url_for("planilla.period_view", period=period))
    msg = f"Planilla cerrada. Se generaron {created} pago(s) de planilla pendientes para el Telecrédito."
    if existing:
        msg += f" {existing} ya existían y no se duplicaron."
    flash(msg, "success")
    return back


@bp.route("/periodo/<period>/reabrir", methods=["POST"])
@permission_required("pagos_personal", "edit")
def period_reopen(period):
    if not validate_csrf():
        abort(400)
    p = _period_or_404(period)
    back = redirect(url_for("planilla.period_view", period=period))
    if p["status"] != "CERRADO":
        return back
    later = query_one("SELECT period FROM payroll_periods WHERE period > ? AND status = 'CERRADO' LIMIT 1", (period,))
    if later:
        flash(f"No se puede reabrir: la planilla de {period_label(later['period'])} ya está cerrada. Reábrela primero.", "error")
        return back
    execute("DELETE FROM staff_loan_payments WHERE period_id = ?", (p["id"],))
    execute("UPDATE payroll_periods SET status = 'ABIERTO', closed_at = NULL, closed_by = NULL WHERE id = ?", (p["id"],))
    log_activity("pagos_personal", "REABRIR", f"Planilla {period_label(period)}: reabierta", entity_type="planilla", entity_id=p["id"],
                 entity_url=url_for("planilla.period_view", period=period))
    flash("Planilla reabierta. Los pagos de planilla que se generaron al cerrar siguen en Pagos: corrígelos o elimínalos allí si cambian los montos.", "info")
    return back


@bp.route("/periodo/<period>/eliminar", methods=["POST"])
@permission_required("pagos_personal", "edit")
def period_delete(period):
    if not validate_csrf():
        abort(400)
    p = _period_or_404(period)
    if p["status"] != "ABIERTO":
        flash("Solo se puede eliminar una planilla abierta.", "error")
        return redirect(url_for("planilla.period_view", period=period))
    execute("DELETE FROM payroll_lines WHERE period_id = ?", (p["id"],))
    execute("DELETE FROM payroll_items WHERE period_id = ?", (p["id"],))
    execute("DELETE FROM payroll_periods WHERE id = ?", (p["id"],))
    log_activity("pagos_personal", "ELIMINAR", f"Planilla {period_label(period)}: periodo eliminado")
    flash("Periodo eliminado.", "success")
    return redirect(url_for("planilla.index"))


@bp.route("/periodo/<period>/conceptos", methods=["POST"])
@permission_required("pagos_personal", "edit")
def item_add(period):
    if not validate_csrf():
        abort(400)
    p = _period_or_404(period)
    back = redirect(url_for("planilla.period_view", period=period))
    if p["status"] != "ABIERTO":
        flash("La planilla está cerrada.", "error")
        return back
    f = request.form
    staff = query_one("SELECT * FROM staff WHERE id = ?", (f.get("staff_id", type=int),))
    kind = f.get("kind")
    concept = (f.get("concept") or "").strip()
    amount = parse_float(f.get("amount"), None)
    if staff is None or kind not in ("INGRESO", "DESCUENTO") or not concept or amount is None or amount <= 0:
        flash("Elige la persona, el tipo, escribe el concepto y un monto mayor a cero.", "error")
        return back
    user = getattr(g, "user", None)
    execute(
        """INSERT INTO payroll_items (period_id, staff_id, kind, concept, amount, taxable, created_by)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (p["id"], staff["id"], kind, concept, amount, 0 if (kind == "INGRESO" and f.get("nontaxable")) else 1, user["id"] if user else None),
    )
    log_activity("pagos_personal", "CREAR", f"Planilla {period_label(period)} — {'ingreso' if kind == 'INGRESO' else 'descuento'} {concept} S/ {amount:.2f}: {staff['name']}",
                 entity_type="planilla", entity_id=p["id"], entity_url=url_for("planilla.period_view", period=period))
    flash("Concepto agregado. Vuelve a calcular para ver el efecto.", "success")
    return back


@bp.route("/periodo/<period>/conceptos/<int:item_id>/eliminar", methods=["POST"])
@permission_required("pagos_personal", "edit")
def item_delete(period, item_id):
    if not validate_csrf():
        abort(400)
    p = _period_or_404(period)
    row = query_one("SELECT * FROM payroll_items WHERE id = ? AND period_id = ?", (item_id, p["id"]))
    if row is None:
        abort(404)
    if p["status"] != "ABIERTO":
        flash("La planilla está cerrada.", "error")
    else:
        execute("DELETE FROM payroll_items WHERE id = ?", (item_id,))
        flash("Concepto eliminado. Vuelve a calcular.", "success")
    return redirect(url_for("planilla.period_view", period=period))


def _boleta_context(p, line):
    staff = query_one("SELECT * FROM staff WHERE id = ?", (line["staff_id"],))
    items = query_all("SELECT * FROM payroll_items WHERE period_id = ? AND staff_id = ? ORDER BY id", (p["id"], line["staff_id"]))
    return {
        "p": p, "l": line, "s": staff, "extra": items, "title": period_label(p["period"]), "regime_labels": REGIME_LABELS,
        "pension_labels": PENSION_LABELS,
    }


@bp.route("/periodo/<period>/boleta/<int:staff_id>")
@permission_required("pagos_personal", "view")
def boleta(period, staff_id):
    p = _period_or_404(period)
    line = query_one(
        "SELECT l.* FROM payroll_lines l WHERE l.period_id = ? AND l.staff_id = ?", (p["id"], staff_id)
    )
    if line is None:
        abort(404)
    return render_template("planilla/boleta.html", boletas=[_boleta_context(p, line)])


@bp.route("/periodo/<period>/boletas")
@permission_required("pagos_personal", "view")
def boletas(period):
    p = _period_or_404(period)
    lines = query_all(
        """SELECT l.* FROM payroll_lines l JOIN staff s ON s.id = l.staff_id WHERE l.period_id = ? ORDER BY s.name""", (p["id"],)
    )
    if not lines:
        abort(404)
    return render_template("planilla/boleta.html", boletas=[_boleta_context(p, l) for l in lines])


@bp.route("/periodo/<period>/excel")
@permission_required("pagos_personal", "view")
def period_excel(period):
    p = _period_or_404(period)
    lines = _lines_for(p["id"])
    wb = Workbook()
    ws = wb.active
    ws.title = period
    heads = [
        ("Persona", "staff_name"), ("DNI", "document_number"), ("Régimen", "regime"), ("Días", "days_worked"),
        ("Sueldo básico", "basic_salary"), ("Sueldo ganado", "salary_earned"), ("Asig. familiar", "family_allowance"),
        ("Otros ingresos", "other_income"), ("Otros ingresos no afectos", "other_income_nontaxable"),
        ("Gratificación", "gratification"), ("Bonif. 9%", "gratification_bonus"), ("Total ingresos", "gross_total"),
        ("Faltas", "absence_deduction"), ("Tardanzas", "tardiness_deduction"), ("Subsidio EsSalud", "subsidy_deduction"),
        ("Pensión (ONP/AFP)", "pension_deduction"), ("5ta categoría", "fifth_deduction"), ("Préstamos", "loan_deduction"),
        ("Adelantos", "advance_deduction"), ("Deuda reconocida", "debt_deduction"), ("Otros descuentos", "other_deduction"),
        ("Total descuentos", "total_deductions"), ("NETO A PAGAR", "net_pay"), ("EsSalud empleador", "essalud_employer"),
        ("Costo empresa", "employer_cost"),
    ]
    ws["A1"] = f"Planilla {period_label(period)} ({p['status'].lower()})"
    ws["A1"].font = Font(bold=True, size=14, color="1D4ED8")
    for c, (h, _k) in enumerate(heads, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, l in enumerate(lines, start=4):
        for c, (_h, k) in enumerate(heads, start=1):
            ws.cell(row=r, column=c, value=REGIME_LABELS.get(l[k], l[k]) if k == "regime" else l[k])
    tr = 4 + len(lines)
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
    for c, (_h, k) in enumerate(heads, start=1):
        if c >= 5 and k != "basic_salary":
            col = ws.cell(row=3, column=c).column_letter
            ws.cell(row=tr, column=c, value=f"=SUM({col}4:{col}{tr - 1})").font = Font(bold=True)
    ws.column_dimensions["A"].width = 30
    for c in range(2, len(heads) + 1):
        ws.column_dimensions[ws.cell(row=3, column=c).column_letter].width = 14
    ws.freeze_panes = "B4"
    buf = io.BytesIO()
    wb.save(buf)
    log_activity("pagos_personal", "EXPORTAR", f"Planilla {period_label(period)}: Excel", entity_type="planilla", entity_id=p["id"])
    return Response(buf.getvalue(), mimetype=XLSX_MIME, headers={"Content-Disposition": f'attachment; filename="planilla_{period}.xlsx"'})


# --- Préstamos, adelantos y reconocimientos de deuda -----------------------

@bp.route("/prestamos")
@permission_required("pagos_personal", "view")
def loans():
    show = request.args.get("ver", "")
    rows = query_all(
        "SELECT l.*, s.name AS staff_name FROM staff_loans l JOIN staff s ON s.id = l.staff_id ORDER BY l.id DESC"
    )
    out = []
    for r in rows:
        remaining, paid = loan_remaining(r["id"], r["amount"])
        state = "ANULADO" if r["status"] == "ANULADO" else ("PAGADO" if remaining <= 0.005 else "ACTIVO")
        if show != "todos" and state != "ACTIVO":
            continue
        out.append({"row": r, "remaining": remaining, "paid": paid, "state": state})
    people = query_all("SELECT id, name FROM staff WHERE status = 'ACTIVO' ORDER BY name")
    return render_template(
        "planilla/prestamos.html", loans=out, show=show, people=people, kinds=LOAN_KINDS, kind_labels=LOAN_KIND_LABELS,
        default_period=today_str()[:7], label=period_label,
    )


@bp.route("/prestamos/nuevo", methods=["POST"])
@permission_required("pagos_personal", "edit")
def loan_add():
    if not validate_csrf():
        abort(400)
    f = request.form
    staff = query_one("SELECT * FROM staff WHERE id = ?", (f.get("staff_id", type=int),))
    kind = f.get("kind")
    amount = parse_float(f.get("amount"), None)
    installments = f.get("installments", type=int) or 1
    start = (f.get("start_period") or "").strip()
    error = None
    if staff is None:
        error = "Elige a la persona."
    elif kind not in LOAN_KIND_LABELS:
        error = "Elige el tipo."
    elif amount is None or amount <= 0:
        error = "Escribe un monto mayor a cero."
    elif not _PERIOD_RE.match(start):
        error = "Elige el mes desde el que se descuenta."
    elif installments < 1 or installments > 60:
        error = "Las cuotas deben estar entre 1 y 60."
    if kind == "ADELANTO":
        installments = 1
    filename = None
    file_storage = request.files.get("loan_file")
    if not error and file_storage and file_storage.filename:
        filename = _save_binary_attachment(file_storage, storage.save_staff_document)
        if not filename:
            error = "El archivo no es válido: sube una foto o un PDF."
    if error:
        flash(error, "error")
        return redirect(url_for("planilla.loans"))
    per = round(-(-amount * 100 // installments) / 100, 2)  # cuota redondeada hacia arriba; la última se ajusta sola
    user = getattr(g, "user", None)
    loan_id = execute(
        """INSERT INTO staff_loans (staff_id, kind, concept, amount, installments, installment_amount, start_period, notes,
           filename, created_by) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (staff["id"], kind, (f.get("concept") or "").strip() or None, amount, installments, per, start,
         (f.get("notes") or "").strip() or None, filename, user["id"] if user else None),
    )
    log_activity("pagos_personal", "CREAR", f"{LOAN_KIND_LABELS[kind]} S/ {amount:.2f} en {installments} cuota(s): {staff['name']}",
                 entity_type="empleado", entity_id=staff["id"], entity_url=url_for("planilla.loans"))
    flash(f"{LOAN_KIND_LABELS[kind]} registrado: se descuenta desde {period_label(start)}.", "success")
    return redirect(url_for("planilla.loans"))


@bp.route("/prestamos/<int:loan_id>/anular", methods=["POST"])
@permission_required("pagos_personal", "edit")
def loan_cancel(loan_id):
    if not validate_csrf():
        abort(400)
    row = query_one("SELECT l.*, s.name AS staff_name FROM staff_loans l JOIN staff s ON s.id = l.staff_id WHERE l.id = ?", (loan_id,))
    if row is None:
        abort(404)
    if row["status"] == "ACTIVO":
        execute("UPDATE staff_loans SET status = 'ANULADO' WHERE id = ?", (loan_id,))
        log_activity("pagos_personal", "ANULAR", f"{LOAN_KIND_LABELS.get(row['kind'], row['kind'])} S/ {row['amount']:.2f}: {row['staff_name']}: anulado",
                     entity_type="empleado", entity_id=row["staff_id"])
        flash("Anulado: ya no se descuenta en las próximas planillas (lo ya descontado queda registrado).", "success")
    return redirect(url_for("planilla.loans"))


@bp.route("/prestamos/<int:loan_id>/archivo")
@permission_required("pagos_personal", "view")
def loan_file(loan_id):
    row = query_one("SELECT filename FROM staff_loans WHERE id = ?", (loan_id,))
    if row is None or not row["filename"]:
        abort(404)
    if storage.using_s3():
        return redirect(storage.staff_document_url(row["filename"]))
    return send_from_directory(storage.local_staff_documents_dir(), row["filename"])


# --- CTS --------------------------------------------------------------------

@bp.route("/cts")
@permission_required("pagos_personal", "view")
def cts_view():
    today = date.today()
    sem = request.args.get("semestre") or (f"{today.year}-05" if today.month <= 8 else f"{today.year}-11")
    if not re.match(r"^\d{4}-(05|11)$", sem):
        abort(404)
    year, dm = int(sem[:4]), int(sem[5:7])
    params = get_params()
    rows, total = [], 0.0
    for st in query_all("SELECT * FROM staff WHERE in_payroll = 1 ORDER BY name"):
        res = cts_for(st, dm, year, params, effective_salary)
        if res:
            rows.append({"staff": st, **res})
            total += res["cts"]
    if request.args.get("formato") == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "CTS"
        heads = ["Persona", "DNI", "Régimen", "Desde", "Hasta", "Meses", "Días", "Sueldo básico", "Asig. familiar", "Gratificación (base)", "Remuneración computable", "CTS a depositar"]
        ws["A1"] = f"CTS {'noviembre ' + str(year - 1) + ' – abril ' + str(year) if dm == 5 else 'mayo – octubre ' + str(year)} (depósito 15 de {'mayo' if dm == 5 else 'noviembre'} {year})"
        ws["A1"].font = Font(bold=True, size=13, color="1D4ED8")
        for c, h in enumerate(heads, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D4ED8")
        for r, x in enumerate(rows, start=4):
            vals = [x["staff"]["name"], x["staff"]["document_number"], REGIME_LABELS.get(x["regime"]), x["start"], x["end"], x["months"], x["days"],
                    x["basic"], x["family"], x["grati"], x["computable"], x["cts"]]
            for c, v in enumerate(vals, start=1):
                ws.cell(row=r, column=c, value=v)
        ws.cell(row=4 + len(rows), column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=4 + len(rows), column=12, value=f"=SUM(L4:L{3 + len(rows)})").font = Font(bold=True)
        for c in range(1, 13):
            ws.column_dimensions[ws.cell(row=3, column=c).column_letter].width = 18 if c > 1 else 30
        buf = io.BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(), mimetype=XLSX_MIME, headers={"Content-Disposition": f'attachment; filename="cts_{sem}.xlsx"'})
    options = [f"{y}-{m}" for y in (year - 1, year, year + 1) for m in ("05", "11")]
    return render_template("planilla/cts.html", sem=sem, year=year, dm=dm, rows=rows, total=r2(total), options=options)
@bp.route("/parametros", methods=["GET", "POST"])
@permission_required("pagos_personal", "edit")
def params_view():
    if request.method == "POST":
        if not validate_csrf():
            abort(400)
        user = getattr(g, "user", None)
        changed, bad = [], []
        for row in get_param_rows():
            raw = request.form.get(f"p_{row['key']}")
            if raw is None or raw.strip() == "":
                continue
            value = parse_float(raw, None)
            if value is None or value < 0:
                bad.append(row["label"])
                continue
            if abs(value - row["value"]) > 1e-9:
                set_param(row["key"], value, user["id"] if user else None)
                changed.append(f"{row['label']}: {row['value']:g} → {value:g}")
        if bad:
            flash("Valor inválido en: " + "; ".join(bad) + ". No se guardó ese cambio.", "error")
        if changed:
            log_activity("pagos_personal", "EDITAR", "Parámetros de planilla: " + "; ".join(changed))
            flash(f"Parámetros guardados ({len(changed)} cambio(s)).", "success")
        elif not bad:
            flash("No hubo cambios.", "info")
        return redirect(url_for("planilla.params_view"))
    return render_template("planilla/parametros.html", rows=get_param_rows())


@bp.route("/personal")
@permission_required("pagos_personal", "view")
def profiles():
    show = request.args.get("ver", "")
    rows = query_all("SELECT * FROM staff WHERE status = 'ACTIVO' ORDER BY in_payroll DESC, name")
    if show == "planilla":
        rows = [r for r in rows if r["in_payroll"]]
    return render_template(
        "planilla/personal.html", rows=rows, show=show, regime_labels=REGIME_LABELS, pension_labels=PENSION_LABELS,
        afp_labels=AFP_LABELS, salary_of=effective_salary,
    )


@bp.route("/personal/<int:staff_id>")
@permission_required("pagos_personal", "view")
def profile(staff_id):
    staff = _get_staff_or_404(staff_id)
    contract = current_contract(staff_id)
    return render_template(
        "planilla/perfil.html", staff=staff, regimes=REGIMES, pension_systems=PENSION_SYSTEMS, afp_names=AFP_NAMES,
        commission_types=COMMISSION_TYPES, contract=contract, salary=effective_salary(staff),
        regime_vacation_days=REGIME_VACATION_DAYS,
    )


@bp.route("/personal/<int:staff_id>/guardar", methods=["POST"])
@permission_required("pagos_personal", "edit")
def profile_save(staff_id):
    staff = _get_staff_or_404(staff_id)
    if not validate_csrf():
        abort(400)
    f = request.form
    regime = f.get("labor_regime", "GENERAL")
    pension = f.get("pension_system", "ONP")
    afp = f.get("afp_name") or None
    commission = f.get("afp_commission_type", "FLUJO")
    errors = []
    if regime not in REGIME_LABELS:
        errors.append("Elige el régimen laboral.")
    if pension not in PENSION_LABELS:
        errors.append("Elige el sistema de pensiones.")
    if pension == "AFP" and afp not in AFP_LABELS:
        errors.append("Elige la AFP.")
    if commission not in dict(COMMISSION_TYPES):
        commission = "FLUJO"
    salary_raw = (f.get("basic_salary") or "").strip()
    salary = parse_float(salary_raw, None) if salary_raw else None
    if salary_raw and (salary is None or salary < 0):
        errors.append("El sueldo básico no es válido.")
    prior_income = parse_float(f.get("fifth_prior_income"), 0) or 0
    prior_withheld = parse_float(f.get("fifth_prior_withheld"), 0) or 0
    if prior_income < 0 or prior_withheld < 0:
        errors.append("Los montos de 5ta categoría no pueden ser negativos.")
    if errors:
        for e in errors:
            flash(e, "error")
        return redirect(url_for("planilla.profile", staff_id=staff_id))
    vac_days = staff["vacation_days_per_year"]
    if regime != staff["labor_regime"]:
        vac_days = REGIME_VACATION_DAYS[regime]
    execute(
        """UPDATE staff SET in_payroll = ?, labor_regime = ?, pension_system = ?, afp_name = ?,
           afp_commission_type = ?, cuspp = ?, basic_salary = ?, family_allowance = ?, fifth_prior_income = ?,
           fifth_prior_withheld = ?, vacation_days_per_year = ? WHERE id = ?""",
        (
            1 if f.get("in_payroll") else 0, regime, pension, afp if pension == "AFP" else None, commission,
            (f.get("cuspp") or "").strip() or None, salary, 1 if f.get("family_allowance") else 0,
            prior_income, prior_withheld, vac_days, staff_id,
        ),
    )
    log_activity(
        "pagos_personal", "EDITAR", f"Perfil de planilla: {staff['name']} ({REGIME_LABELS[regime]}, {PENSION_LABELS[pension]})",
        entity_type="empleado", entity_id=staff_id, entity_url=url_for("planilla.profile", staff_id=staff_id),
    )
    msg = "Perfil de planilla guardado."
    if regime != staff["labor_regime"]:
        msg += f" Días de vacaciones por año: {vac_days:g}."
    flash(msg, "success")
    return redirect(url_for("planilla.profile", staff_id=staff_id))
