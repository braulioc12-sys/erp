"""Control de gastos (9 oct, pedido de Braulio: "un módulo pero solo para mi
usuario que se llamará control de gastos... se registra un gasto y luego se
enlaza con un abono").

Reemplaza las hojas "pagos 2026" y "abonos 2026" de su Excel de deuda:

- GASTO: algo que Braulio pagó (tarjeta, transferencia, yape...) por una de
  las empresas. Puede estar en dólares (monto USD × tipo de cambio).
- ABONO: lo que la empresa le devuelve. Un abono cubre uno o varios gastos
  (`cg_gastos.abono_id`, el equivalente a la columna "Nro Op" del Excel).
- La deuda es gastos - abonos. Un abono "cuadra" cuando la suma de sus
  gastos enlazados es igual a su monto.

ACCESO: exclusivo por correo de login (`CONTROL_GASTOS_EMAILS` en config.py),
NO por rol: ni siquiera otros administradores lo ven. Para quien no esté en
la lista el módulo responde 404 (no revela que existe). Tampoco escribe en
el registro de Actividad (activity_log), para que no aparezca ahí.
"""
import functools
import io
import re
from datetime import date, datetime

from flask import (
    Blueprint, Response, abort, current_app, flash, g, redirect, render_template, request, url_for,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.auth import validate_csrf
from app.db import execute, query_all, query_one
from app.helpers import parse_date, parse_float, today_str

bp = Blueprint("control_gastos", __name__, url_prefix="/control-gastos")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PAY_METHODS = ["tarjeta", "transferencia", "yape", "efectivo", "otro"]
MONTH_NAMES = [
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
]
# Diferencia máxima (S/) para considerar que un abono "cuadra": en el Excel
# original hay abonos con centavos de diferencia por redondeo.
TOLERANCE = 0.10
PAGE_SIZE = 200

STATE_LABELS = {
    "SIN_ENLAZAR": "Sin gastos",
    "CUADRA": "Cuadra",
    "FALTAN": "Faltan gastos",
    "EXCEDE": "Gastos exceden",
}


# --- acceso -----------------------------------------------------------------

def allowed_emails():
    raw = current_app.config.get("CONTROL_GASTOS_EMAILS", "") or ""
    return {e.strip().lower() for e in raw.split(",") if e.strip()}


def user_can_access(user=None):
    user = user if user is not None else g.get("user")
    if not user:
        return False
    return (user["email"] or "").strip().lower() in allowed_emails()


def control_gastos_required(view):
    @functools.wraps(view)
    def wrapped(**kwargs):
        if g.user is None:
            return redirect(url_for("auth.login", next=request.path))
        if not user_can_access(g.user):
            abort(404)
        return view(**kwargs)

    return wrapped


# --- utilidades ---------------------------------------------------------------

def _norm_company(value):
    return re.sub(r"\s+", " ", (value or "").strip().lower())


def _norm_code(value):
    return re.sub(r"\s+", "", (value or "").strip().lower())


def _round(value):
    return round(float(value or 0), 2)


def abono_state(amount, linked, n_gastos):
    """(clave, diferencia): diferencia = monto del abono - gastos enlazados."""
    diff = _round((amount or 0) - (linked or 0))
    if not n_gastos:
        return "SIN_ENLAZAR", diff
    if abs(diff) <= TOLERANCE:
        return "CUADRA", diff
    return ("FALTAN" if diff > 0 else "EXCEDE"), diff


_ABONO_SQL = """
    SELECT a.*, COALESCE(s.linked, 0) AS linked, COALESCE(s.n, 0) AS n_gastos
    FROM cg_abonos a
    LEFT JOIN (SELECT abono_id, SUM(amount) AS linked, COUNT(*) AS n
               FROM cg_gastos WHERE abono_id IS NOT NULL GROUP BY abono_id) s ON s.abono_id = a.id
"""


def _abonos_with_state(where="", params=()):
    rows = []
    for a in query_all(f"{_ABONO_SQL} {where} ORDER BY a.abono_date DESC, a.id DESC", params):
        d = dict(a)
        d["state"], d["diff"] = abono_state(a["amount"], a["linked"], a["n_gastos"])
        d["state_label"] = STATE_LABELS[d["state"]]
        rows.append(d)
    return rows


def _next_abono_code():
    best = 0
    for r in query_all("SELECT code FROM cg_abonos"):
        m = re.fullmatch(r"a(\d+)", r["code"] or "")
        if m:
            best = max(best, int(m.group(1)))
    return f"a{best + 1}"


def _companies():
    seen = []
    for table in ("cg_gastos", "cg_abonos"):
        for r in query_all(f"SELECT DISTINCT company FROM {table} WHERE company IS NOT NULL AND company <> ''"):
            if r["company"] not in seen:
                seen.append(r["company"])
    return sorted(seen)


def _years():
    years = {int(today_str()[:4])}
    for sql in ("SELECT DISTINCT substr(expense_date,1,4) y FROM cg_gastos",
                "SELECT DISTINCT substr(abono_date,1,4) y FROM cg_abonos"):
        for r in query_all(sql):
            if r["y"] and r["y"].isdigit():
                years.add(int(r["y"]))
    return sorted(years, reverse=True)


def _totals(company=""):
    """Totales generales (todo el tiempo) por empresa o de una sola."""
    params = [company] if company else []
    cond = " WHERE company = ?" if company else ""
    g_row = query_one(f"SELECT COALESCE(SUM(amount),0) t, COUNT(*) n FROM cg_gastos{cond}", params)
    a_row = query_one(f"SELECT COALESCE(SUM(amount),0) t, COUNT(*) n FROM cg_abonos{cond}", params)
    pend_cond = (" AND company = ?" if company else "")
    p_row = query_one(
        f"SELECT COALESCE(SUM(amount),0) t, COUNT(*) n FROM cg_gastos WHERE abono_id IS NULL{pend_cond}", params
    )
    return {
        "gastos": _round(g_row["t"]), "n_gastos": g_row["n"],
        "abonos": _round(a_row["t"]), "n_abonos": a_row["n"],
        "deuda": _round(g_row["t"] - a_row["t"]),
        "sin_abono": _round(p_row["t"]), "n_sin_abono": p_row["n"],
    }


def _fx_amount(form):
    """(monto_soles, usd, tipo_cambio, error) desde el formulario de gasto:
    si no se indica el monto en soles pero sí USD y tipo de cambio, se calcula."""
    amount = parse_float(form.get("amount"), None) if (form.get("amount") or "").strip() else None
    usd = parse_float(form.get("foreign_amount"), None) if (form.get("foreign_amount") or "").strip() else None
    rate = parse_float(form.get("exchange_rate"), None) if (form.get("exchange_rate") or "").strip() else None
    if usd is not None and rate is not None and amount is None:
        amount = _round(usd * rate)
    if amount is None or amount <= 0:
        return None, usd, rate, "Indica el monto del gasto en soles (o el monto en dólares con su tipo de cambio)."
    if (usd is None) != (rate is None):
        usd, rate = None, None  # un dato suelto no sirve: se ignora
    return _round(amount), usd, rate, None


def _page_back(default_endpoint="control_gastos.index", **values):
    nxt = request.form.get("next") or request.args.get("next")
    if nxt and nxt.startswith("/control-gastos"):
        return redirect(nxt)
    return redirect(url_for(default_endpoint, **values))


# --- gastos ------------------------------------------------------------------

def _gasto_filters():
    return {
        "q": request.args.get("q", "").strip(),
        "company": _norm_company(request.args.get("empresa")),
        "estado": request.args.get("estado", ""),
        "month": request.args.get("mes", "").strip(),
        "method": request.args.get("medio", "").strip().lower(),
    }


def _gasto_where(f):
    where, params = [], []
    if f["q"]:
        where.append("(lower(g.supplier) LIKE ? OR lower(COALESCE(g.notes,'')) LIKE ? OR lower(COALESCE(a.code,'')) LIKE ?)")
        like = f"%{f['q'].lower()}%"
        params += [like, like, like]
    if f["company"]:
        where.append("g.company = ?")
        params.append(f["company"])
    if f["estado"] == "sin_abono":
        where.append("g.abono_id IS NULL")
    elif f["estado"] == "con_abono":
        where.append("g.abono_id IS NOT NULL")
    if re.fullmatch(r"\d{4}-\d{2}", f["month"]):
        where.append("substr(g.expense_date,1,7) = ?")
        params.append(f["month"])
    elif re.fullmatch(r"\d{4}", f["month"]):
        where.append("substr(g.expense_date,1,4) = ?")
        params.append(f["month"])
    if f["method"]:
        where.append("lower(COALESCE(g.pay_method,'')) = ?")
        params.append(f["method"])
    return (" WHERE " + " AND ".join(where)) if where else "", params


_GASTO_SELECT = """SELECT g.*, a.code AS abono_code
                   FROM cg_gastos g LEFT JOIN cg_abonos a ON a.id = g.abono_id"""


@bp.route("")
@control_gastos_required
def index():
    f = _gasto_filters()
    where, params = _gasto_where(f)
    page = max(request.args.get("pag", 1, type=int) or 1, 1)
    summary = query_one(
        f"SELECT COALESCE(SUM(g.amount),0) t, COUNT(*) n FROM cg_gastos g LEFT JOIN cg_abonos a ON a.id = g.abono_id{where}",
        params,
    )
    rows = query_all(
        f"{_GASTO_SELECT}{where} ORDER BY g.expense_date DESC, g.id DESC LIMIT ? OFFSET ?",
        params + [PAGE_SIZE, (page - 1) * PAGE_SIZE],
    )
    open_abonos = _abonos_with_state()
    qs = {k: v for k, v in request.args.items() if k != "pag" and v}
    return render_template(
        "control_gastos/index.html", rows=rows, f=f, page=page, page_size=PAGE_SIZE,
        total_rows=summary["n"], total_amount=_round(summary["t"]), companies=_companies(),
        methods=PAY_METHODS, totals=_totals(), totals_by_company={c: _totals(c) for c in _companies()},
        abonos=open_abonos, next_code=_next_abono_code(), today=today_str(), qs=qs,
        current_url=request.full_path.rstrip("?"),
    )


def _gasto_form_context(gasto=None):
    return dict(
        gasto=gasto, companies=_companies(), methods=PAY_METHODS, today=today_str(),
        abonos=_abonos_with_state(),
    )


def _read_gasto_form():
    """(datos, error)"""
    form = request.form
    supplier = (form.get("supplier") or "").strip()
    expense_date = parse_date(form.get("expense_date"))
    amount, usd, rate, error = _fx_amount(form)
    if not supplier:
        error = "Indica el proveedor o a quién se pagó."
    elif not expense_date:
        error = "Indica la fecha del gasto."
    if error:
        return None, error
    abono_id = form.get("abono_id", type=int)
    if abono_id and query_one("SELECT id FROM cg_abonos WHERE id = ?", (abono_id,)) is None:
        abono_id = None
    return {
        "expense_date": expense_date,
        "pay_method": (form.get("pay_method") or "").strip().lower(),
        "supplier": supplier,
        "amount": amount, "foreign_amount": usd, "exchange_rate": rate,
        "company": _norm_company(form.get("company")),
        "notes": (form.get("notes") or "").strip(),
        "abono_id": abono_id,
    }, None


@bp.route("/gastos/nuevo", methods=["GET", "POST"])
@control_gastos_required
def gasto_new():
    if request.method == "POST":
        if not validate_csrf():
            abort(400)
        data, error = _read_gasto_form()
        if error:
            flash(error, "error")
            return render_template("control_gastos/gasto_form.html", **_gasto_form_context(request.form))
        execute(
            """INSERT INTO cg_gastos (expense_date, pay_method, supplier, amount, foreign_amount, exchange_rate,
                                      company, notes, abono_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (data["expense_date"], data["pay_method"], data["supplier"], data["amount"], data["foreign_amount"],
             data["exchange_rate"], data["company"], data["notes"], data["abono_id"]),
        )
        flash(f'Gasto registrado: {data["supplier"]} por S/ {data["amount"]:,.2f}.', "success")
        if request.form.get("seguir") == "1":
            return redirect(url_for("control_gastos.gasto_new"))
        return redirect(url_for("control_gastos.index"))
    return render_template("control_gastos/gasto_form.html", **_gasto_form_context())


@bp.route("/gastos/<int:gasto_id>/editar", methods=["GET", "POST"])
@control_gastos_required
def gasto_edit(gasto_id):
    gasto = query_one("SELECT * FROM cg_gastos WHERE id = ?", (gasto_id,))
    if gasto is None:
        abort(404)
    if request.method == "POST":
        if not validate_csrf():
            abort(400)
        data, error = _read_gasto_form()
        if error:
            flash(error, "error")
            ctx = _gasto_form_context(dict(request.form, id=gasto_id))
            return render_template("control_gastos/gasto_form.html", **ctx)
        execute(
            """UPDATE cg_gastos SET expense_date = ?, pay_method = ?, supplier = ?, amount = ?, foreign_amount = ?,
                      exchange_rate = ?, company = ?, notes = ?, abono_id = ? WHERE id = ?""",
            (data["expense_date"], data["pay_method"], data["supplier"], data["amount"], data["foreign_amount"],
             data["exchange_rate"], data["company"], data["notes"], data["abono_id"], gasto_id),
        )
        flash("Gasto actualizado.", "success")
        return _page_back()
    return render_template("control_gastos/gasto_form.html", **_gasto_form_context(gasto))


@bp.route("/gastos/<int:gasto_id>/eliminar", methods=["POST"])
@control_gastos_required
def gasto_delete(gasto_id):
    if not validate_csrf():
        abort(400)
    gasto = query_one("SELECT * FROM cg_gastos WHERE id = ?", (gasto_id,))
    if gasto is None:
        abort(404)
    execute("DELETE FROM cg_gastos WHERE id = ?", (gasto_id,))
    flash(f'Gasto "{gasto["supplier"]}" eliminado.', "success")
    return _page_back()


@bp.route("/gastos/<int:gasto_id>/desenlazar", methods=["POST"])
@control_gastos_required
def gasto_unlink(gasto_id):
    if not validate_csrf():
        abort(400)
    gasto = query_one("SELECT * FROM cg_gastos WHERE id = ?", (gasto_id,))
    if gasto is None:
        abort(404)
    execute("UPDATE cg_gastos SET abono_id = NULL WHERE id = ?", (gasto_id,))
    flash("El gasto quedó sin abono.", "success")
    return _page_back()


@bp.route("/enlazar", methods=["POST"])
@control_gastos_required
def link():
    """Enlaza los gastos marcados con un abono existente o con uno nuevo que se
    crea en el mismo paso (por defecto con la suma de los gastos)."""
    if not validate_csrf():
        abort(400)
    ids = [i for i in request.form.getlist("gasto_ids", type=int) if i]
    if not ids:
        flash("Marca al menos un gasto para enlazar.", "error")
        return _page_back()
    marks = ",".join("?" for _ in ids)
    gastos = query_all(f"SELECT * FROM cg_gastos WHERE id IN ({marks})", ids)
    free = [g_ for g_ in gastos if g_["abono_id"] is None]
    if not free:
        flash("Los gastos marcados ya tienen abono.", "error")
        return _page_back()

    if request.form.get("modo") == "nuevo":
        code = _norm_code(request.form.get("code")) or _next_abono_code()
        if query_one("SELECT id FROM cg_abonos WHERE code = ?", (code,)):
            flash(f'Ya existe el abono "{code}". Elige otro código o enlázalos a ese abono.', "error")
            return _page_back()
        total = _round(sum(g_["amount"] for g_ in free))
        raw_amount = (request.form.get("amount") or "").strip()
        amount = _round(parse_float(raw_amount, total)) if raw_amount else total
        companies = {g_["company"] for g_ in free if g_["company"]}
        company = _norm_company(request.form.get("company")) or (companies.pop() if len(companies) == 1 else "")
        abono_id = execute(
            "INSERT INTO cg_abonos (code, abono_date, amount, company, description) VALUES (?, ?, ?, ?, ?)",
            (code, parse_date(request.form.get("abono_date")) or today_str(), amount, company,
             (request.form.get("description") or "").strip()),
        )
    else:
        abono_id = request.form.get("abono_id", type=int)
        abono = query_one("SELECT * FROM cg_abonos WHERE id = ?", (abono_id,)) if abono_id else None
        if abono is None:
            flash("Elige el abono con el que se enlazan los gastos.", "error")
            return _page_back()

    marks = ",".join("?" for _ in free)
    execute(f"UPDATE cg_gastos SET abono_id = ? WHERE id IN ({marks})", [abono_id] + [g_["id"] for g_ in free])
    skipped = len(gastos) - len(free)
    msg = f"{len(free)} gasto(s) enlazado(s)."
    if skipped:
        msg += f" {skipped} ya tenían abono y no se tocaron."
    flash(msg, "success")
    return redirect(url_for("control_gastos.abono_detail", abono_id=abono_id))


# --- abonos --------------------------------------------------------------------

@bp.route("/abonos")
@control_gastos_required
def abonos():
    company = _norm_company(request.args.get("empresa"))
    estado = request.args.get("estado", "")
    month = request.args.get("mes", "").strip()
    q = request.args.get("q", "").strip().lower()
    where, params = [], []
    if company:
        where.append("a.company = ?")
        params.append(company)
    if re.fullmatch(r"\d{4}-\d{2}", month):
        where.append("substr(a.abono_date,1,7) = ?")
        params.append(month)
    elif re.fullmatch(r"\d{4}", month):
        where.append("substr(a.abono_date,1,4) = ?")
        params.append(month)
    if q:
        where.append("(lower(a.code) LIKE ? OR lower(COALESCE(a.description,'')) LIKE ?)")
        params += [f"%{q}%", f"%{q}%"]
    rows = _abonos_with_state((" WHERE " + " AND ".join(where)) if where else "", params)
    if estado in STATE_LABELS:
        rows = [r for r in rows if r["state"] == estado]
    return render_template(
        "control_gastos/abonos.html", rows=rows, company=company, estado=estado, month=month, q=q,
        companies=_companies(), state_labels=STATE_LABELS, total=_round(sum(r["amount"] for r in rows)),
        totals=_totals(),
    )


def _read_abono_form(abono_id=None):
    form = request.form
    code = _norm_code(form.get("code"))
    amount = parse_float(form.get("amount"), None) if (form.get("amount") or "").strip() else None
    if not code:
        return None, "Indica el código del abono (por ejemplo a144)."
    dup = query_one("SELECT id FROM cg_abonos WHERE code = ?", (code,))
    if dup and dup["id"] != abono_id:
        return None, f'Ya existe un abono con el código "{code}".'
    if amount is None or amount < 0:
        return None, "Indica el monto del abono."
    return {
        "code": code, "abono_date": parse_date(form.get("abono_date")) or today_str(), "amount": _round(amount),
        "company": _norm_company(form.get("company")), "description": (form.get("description") or "").strip(),
        "notes": (form.get("notes") or "").strip(),
    }, None


@bp.route("/abonos/nuevo", methods=["GET", "POST"])
@control_gastos_required
def abono_new():
    if request.method == "POST":
        if not validate_csrf():
            abort(400)
        data, error = _read_abono_form()
        if error:
            flash(error, "error")
            return render_template("control_gastos/abono_form.html", abono=request.form, companies=_companies(), today=today_str())
        abono_id = execute(
            "INSERT INTO cg_abonos (code, abono_date, amount, company, description, notes) VALUES (?, ?, ?, ?, ?, ?)",
            (data["code"], data["abono_date"], data["amount"], data["company"], data["description"], data["notes"]),
        )
        flash(f'Abono {data["code"]} registrado. Ahora enlaza los gastos que cubre.', "success")
        return redirect(url_for("control_gastos.abono_detail", abono_id=abono_id))
    return render_template(
        "control_gastos/abono_form.html", abono={"code": _next_abono_code()}, companies=_companies(), today=today_str(),
    )


@bp.route("/abonos/<int:abono_id>/editar", methods=["GET", "POST"])
@control_gastos_required
def abono_edit(abono_id):
    abono = query_one("SELECT * FROM cg_abonos WHERE id = ?", (abono_id,))
    if abono is None:
        abort(404)
    if request.method == "POST":
        if not validate_csrf():
            abort(400)
        data, error = _read_abono_form(abono_id)
        if error:
            flash(error, "error")
            return render_template(
                "control_gastos/abono_form.html", abono=dict(request.form, id=abono_id), companies=_companies(), today=today_str(),
            )
        execute(
            """UPDATE cg_abonos SET code = ?, abono_date = ?, amount = ?, company = ?, description = ?, notes = ?
               WHERE id = ?""",
            (data["code"], data["abono_date"], data["amount"], data["company"], data["description"], data["notes"], abono_id),
        )
        flash("Abono actualizado.", "success")
        return redirect(url_for("control_gastos.abono_detail", abono_id=abono_id))
    return render_template("control_gastos/abono_form.html", abono=abono, companies=_companies(), today=today_str())


@bp.route("/abonos/<int:abono_id>")
@control_gastos_required
def abono_detail(abono_id):
    found = _abonos_with_state("WHERE a.id = ?", (abono_id,))
    if not found:
        abort(404)
    abono = found[0]
    linked = query_all(
        f"{_GASTO_SELECT} WHERE g.abono_id = ? ORDER BY g.expense_date, g.id", (abono_id,)
    )
    all_companies = request.args.get("todas") == "1"
    pend_sql = f"{_GASTO_SELECT} WHERE g.abono_id IS NULL"
    pend_params = []
    if abono["company"] and not all_companies:
        pend_sql += " AND g.company = ?"
        pend_params.append(abono["company"])
    pending = query_all(pend_sql + " ORDER BY g.expense_date DESC, g.id DESC", pend_params)
    return render_template(
        "control_gastos/abono_detail.html", abono=abono, linked=linked, pending=pending,
        all_companies=all_companies, state_labels=STATE_LABELS,
        pending_total=_round(sum(p["amount"] for p in pending)),
    )


@bp.route("/abonos/<int:abono_id>/enlazar", methods=["POST"])
@control_gastos_required
def abono_link(abono_id):
    if not validate_csrf():
        abort(400)
    if query_one("SELECT id FROM cg_abonos WHERE id = ?", (abono_id,)) is None:
        abort(404)
    ids = [i for i in request.form.getlist("gasto_ids", type=int) if i]
    if not ids:
        flash("Marca al menos un gasto para enlazar.", "error")
        return redirect(url_for("control_gastos.abono_detail", abono_id=abono_id))
    marks = ",".join("?" for _ in ids)
    execute(f"UPDATE cg_gastos SET abono_id = ? WHERE abono_id IS NULL AND id IN ({marks})", [abono_id] + ids)
    flash("Gastos enlazados al abono.", "success")
    return redirect(url_for("control_gastos.abono_detail", abono_id=abono_id))


@bp.route("/abonos/<int:abono_id>/ajustar-monto", methods=["POST"])
@control_gastos_required
def abono_adjust(abono_id):
    """Pone como monto del abono la suma de sus gastos enlazados."""
    if not validate_csrf():
        abort(400)
    found = _abonos_with_state("WHERE a.id = ?", (abono_id,))
    if not found:
        abort(404)
    execute("UPDATE cg_abonos SET amount = ? WHERE id = ?", (_round(found[0]["linked"]), abono_id))
    flash("El monto del abono ahora es la suma de sus gastos.", "success")
    return redirect(url_for("control_gastos.abono_detail", abono_id=abono_id))


@bp.route("/abonos/<int:abono_id>/eliminar", methods=["POST"])
@control_gastos_required
def abono_delete(abono_id):
    if not validate_csrf():
        abort(400)
    abono = query_one("SELECT * FROM cg_abonos WHERE id = ?", (abono_id,))
    if abono is None:
        abort(404)
    execute("UPDATE cg_gastos SET abono_id = NULL WHERE abono_id = ?", (abono_id,))
    execute("DELETE FROM cg_abonos WHERE id = ?", (abono_id,))
    flash(f'Abono {abono["code"]} eliminado; sus gastos quedaron sin abono.', "success")
    return redirect(url_for("control_gastos.abonos"))


# --- resumen -------------------------------------------------------------------

def _summary_rows(year, company):
    """Un renglón por mes del año: gastos, abonos y diferencia (gastos - abonos)."""
    g_params, a_params = [str(year)], [str(year)]
    g_cond = a_cond = ""
    if company:
        g_cond = a_cond = " AND company = ?"
        g_params.append(company)
        a_params.append(company)
    g_by = {r["m"]: r for r in query_all(
        f"SELECT substr(expense_date,6,2) m, SUM(amount) t, COUNT(*) n FROM cg_gastos "
        f"WHERE substr(expense_date,1,4) = ?{g_cond} GROUP BY m", g_params)}
    a_by = {r["m"]: r for r in query_all(
        f"SELECT substr(abono_date,6,2) m, SUM(amount) t, COUNT(*) n FROM cg_abonos "
        f"WHERE substr(abono_date,1,4) = ?{a_cond} GROUP BY m", a_params)}
    rows = []
    for m in range(1, 13):
        key = f"{m:02d}"
        g_t = _round(g_by[key]["t"]) if key in g_by else 0.0
        a_t = _round(a_by[key]["t"]) if key in a_by else 0.0
        rows.append({
            "month": m, "name": MONTH_NAMES[m - 1],
            "gastos": g_t, "n_gastos": g_by[key]["n"] if key in g_by else 0,
            "abonos": a_t, "n_abonos": a_by[key]["n"] if key in a_by else 0,
            "diff": _round(g_t - a_t),
        })
    return rows


@bp.route("/resumen")
@control_gastos_required
def resumen():
    year = request.args.get("anio", type=int) or int(today_str()[:4])
    company = _norm_company(request.args.get("empresa"))
    rows = _summary_rows(year, company)
    totals_year = {
        "gastos": _round(sum(r["gastos"] for r in rows)), "abonos": _round(sum(r["abonos"] for r in rows)),
    }
    totals_year["diff"] = _round(totals_year["gastos"] - totals_year["abonos"])
    by_company = []
    for c in _companies():
        t = _totals(c)
        t["company"] = c
        by_company.append(t)
    by_company.append(dict(_totals(), company=""))  # "Todas"
    return render_template(
        "control_gastos/resumen.html", rows=rows, year=year, company=company, years=_years(),
        companies=_companies(), totals_year=totals_year, by_company=by_company,
    )


# --- exportar ------------------------------------------------------------------

def _style_header(ws, row, titles):
    for idx, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=idx, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="center")


def build_workbook(year=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"
    titles = ["N° op (abono)", "Medio", "Fecha", "Proveedor", "Monto S/", "Monto USD", "T.C.", "Empresa", "Estado", "Notas"]
    _style_header(ws, 1, titles)
    sql = f"{_GASTO_SELECT}"
    params = []
    if year:
        sql += " WHERE substr(g.expense_date,1,4) = ?"
        params.append(str(year))
    r = 2
    for g_ in query_all(sql + " ORDER BY g.expense_date, g.id", params):
        ws.append([
            g_["abono_code"] or "", g_["pay_method"] or "", g_["expense_date"], g_["supplier"], g_["amount"],
            g_["foreign_amount"], g_["exchange_rate"], g_["company"] or "",
            "Enlazado" if g_["abono_id"] else "Sin abono", g_["notes"] or "",
        ])
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        r += 1
    for idx, w in enumerate([13, 14, 12, 34, 14, 12, 8, 12, 12, 30], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"

    wa = wb.create_sheet("Abonos")
    _style_header(wa, 1, ["Código", "Fecha", "Monto S/", "Empresa", "Descripción", "N° gastos", "Suma gastos", "Diferencia", "Estado"])
    a_where, a_params = "", []
    if year:
        a_where, a_params = "WHERE substr(a.abono_date,1,4) = ?", [str(year)]
    r = 2
    for a in reversed(_abonos_with_state(a_where, a_params)):
        wa.append([a["code"], a["abono_date"], a["amount"], a["company"] or "", a["description"] or "",
                   a["n_gastos"], _round(a["linked"]), a["diff"], a["state_label"]])
        for col in (3, 7, 8):
            wa.cell(row=r, column=col).number_format = "#,##0.00"
        r += 1
    for idx, w in enumerate([10, 12, 14, 12, 30, 10, 14, 12, 16], start=1):
        wa.column_dimensions[get_column_letter(idx)].width = w
    wa.freeze_panes = "A2"

    wr = wb.create_sheet("Resumen")
    _style_header(wr, 1, ["Empresa", "Gastos S/", "Abonos S/", "Deuda (gastos - abonos)", "Gastos sin abono S/", "N° gastos sin abono"])
    for c in _companies() + [""]:
        t = _totals(c)
        wr.append([c or "TOTAL", t["gastos"], t["abonos"], t["deuda"], t["sin_abono"], t["n_sin_abono"]])
    for row in wr.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = "#,##0.00"
    for idx, w in enumerate([14, 16, 16, 24, 20, 18], start=1):
        wr.column_dimensions[get_column_letter(idx)].width = w
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@bp.route("/exportar")
@control_gastos_required
def export_excel():
    year = request.args.get("anio", type=int)
    buffer = build_workbook(year)
    name = f"control_gastos_{year or 'todo'}_{today_str()}.xlsx"
    return Response(
        buffer.getvalue(), mimetype=XLSX_MIME, headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


# --- importar el Excel "deuda" ------------------------------------------------

_FORMULA_FX = re.compile(r"^=\s*([\d.]+)\s*\*\s*([\d.]+)\s*$")


def _cell_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _cell_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _find_sheet(wb, word, year):
    for name in wb.sheetnames:
        low = name.strip().lower()
        if word in low and str(year) in low:
            return name
    return None


def import_workbook(content, year):
    """Lee las hojas "pagos <año>" y "abonos <año>" del Excel de deuda.
    Devuelve (abonos, gastos, avisos) ya normalizados, sin tocar la base."""
    wb_f = load_workbook(io.BytesIO(content))
    wb_v = load_workbook(io.BytesIO(content), data_only=True)
    pagos_name, abonos_name = _find_sheet(wb_v, "pagos", year), _find_sheet(wb_v, "abonos", year)
    if not pagos_name or not abonos_name:
        raise ValueError(
            f'No encontré las hojas "pagos {year}" y "abonos {year}". Hojas del archivo: ' + ", ".join(wb_v.sheetnames)
        )
    warnings = []

    abonos, seen = [], set()
    ws, wf = wb_v[abonos_name], wb_f[abonos_name]
    header_row = next(
        (r for r in range(1, 10) if str(ws.cell(r, 3).value or "").strip().lower().startswith("nro")), 3
    )
    for r in range(header_row + 1, ws.max_row + 1):
        code = _norm_code(str(ws.cell(r, 3).value or ""))
        amount = _cell_number(ws.cell(r, 5).value)
        if not code:
            continue
        if amount is None:
            warnings.append(f'Abono {code} (fila {r}) sin monto: no se importó.')
            continue
        if code in seen:
            warnings.append(f'Abono {code} repetido (fila {r}): se ignoró el duplicado.')
            continue
        seen.add(code)
        abonos.append({
            "code": code, "abono_date": _cell_date(ws.cell(r, 4).value), "amount": _round(amount),
            "company": _norm_company(str(ws.cell(r, 6).value or "")), "description": str(ws.cell(r, 7).value or "").strip(),
        })

    gastos = []
    ws, wf = wb_v[pagos_name], wb_f[pagos_name]
    for r in range(2, ws.max_row + 1):
        fecha = _cell_date(ws.cell(r, 3).value)
        supplier = str(ws.cell(r, 4).value or "").strip()
        if not fecha and not supplier and ws.cell(r, 5).value is None:
            continue
        raw_formula = wf.cell(r, 5).value
        amount = _cell_number(ws.cell(r, 5).value)
        usd = rate = None
        if isinstance(raw_formula, str):
            m = _FORMULA_FX.match(raw_formula)
            if m:
                usd, rate = float(m.group(1)), float(m.group(2))
                if amount is None:
                    amount = usd * rate
        if not fecha or not supplier or amount is None or amount <= 0:
            warnings.append(f"Fila {r} de {pagos_name} incompleta (fecha, proveedor o monto): no se importó.")
            continue
        gastos.append({
            "op": _norm_code(str(ws.cell(r, 1).value or "")), "pay_method": str(ws.cell(r, 2).value or "").strip().lower(),
            "expense_date": fecha, "supplier": supplier, "amount": _round(amount),
            "foreign_amount": usd, "exchange_rate": rate,
            "company": _norm_company(str(ws.cell(r, 6).value or "")), "row": r,
        })
    return abonos, gastos, warnings


@bp.route("/importar", methods=["GET", "POST"])
@control_gastos_required
def import_excel():
    result = None
    existing = query_one("SELECT (SELECT COUNT(*) FROM cg_gastos) g, (SELECT COUNT(*) FROM cg_abonos) a")
    if request.method == "POST":
        if not validate_csrf():
            abort(400)
        year = request.form.get("anio", type=int) or int(today_str()[:4])
        upload = request.files.get("file")
        replace = request.form.get("reemplazar") == "1"
        if upload is None or not upload.filename.lower().endswith(".xlsx"):
            flash("Sube el archivo Excel (.xlsx).", "error")
        elif (existing["g"] or existing["a"]) and not replace:
            flash("Ya hay gastos o abonos cargados. Marca \"Borrar lo cargado y reemplazar\" si quieres volver a importar.", "error")
        else:
            try:
                abonos_in, gastos_in, warnings = import_workbook(upload.read(), year)
            except Exception as exc:  # archivo ilegible o sin las hojas esperadas
                flash(f"No se pudo leer el archivo: {exc}", "error")
            else:
                if replace:
                    execute("UPDATE cg_gastos SET abono_id = NULL")
                    execute("DELETE FROM cg_gastos")
                    execute("DELETE FROM cg_abonos")
                ids = {}
                for a in abonos_in:
                    ids[a["code"]] = execute(
                        "INSERT INTO cg_abonos (code, abono_date, amount, company, description) VALUES (?, ?, ?, ?, ?)",
                        (a["code"], a["abono_date"], a["amount"], a["company"], a["description"]),
                    )
                linked = unlinked = 0
                unknown = {}
                for g_ in gastos_in:
                    abono_id = ids.get(g_["op"]) if g_["op"] else None
                    notes = ""
                    if g_["op"] and abono_id is None:
                        notes = f'N° op en el Excel: {g_["op"]}'
                        unknown[g_["op"]] = unknown.get(g_["op"], 0) + 1
                    linked += 1 if abono_id else 0
                    unlinked += 0 if abono_id else 1
                    execute(
                        """INSERT INTO cg_gastos (expense_date, pay_method, supplier, amount, foreign_amount, exchange_rate,
                                                  company, notes, abono_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (g_["expense_date"], g_["pay_method"], g_["supplier"], g_["amount"], g_["foreign_amount"],
                         g_["exchange_rate"], g_["company"], notes, abono_id),
                    )
                result = {
                    "abonos": len(abonos_in), "gastos": len(gastos_in), "linked": linked, "unlinked": unlinked,
                    "unknown": unknown, "warnings": warnings,
                }
                existing = query_one("SELECT (SELECT COUNT(*) FROM cg_gastos) g, (SELECT COUNT(*) FROM cg_abonos) a")
    return render_template("control_gastos/import.html", result=result, existing=existing, year=int(today_str()[:4]))
